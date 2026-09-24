import csv
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsupportedFileTypeError(ValueError):
    """Raised for a content_type this parser has no extraction logic for.

    A ValueError subclass (not NotImplementedError): this is a bad-input
    condition, not a missing code path. Temporal activities that wrap
    this parser treat ValueError as non-retryable — an unsupported file
    will never become supported on retry, so retrying it just wastes time
    before reaching the same failure.
    """


def extract_text(path: Path, content_type: str) -> str:
    """Extract plain text from a document on disk.

    A plain function, not a class: path/content_type are only ever used
    once per call, there's no state to hold between calls, and a function
    is simpler to unit-test and to wrap in a Temporal activity.
    """
    match content_type:
        case "pdf":
            reader = PdfReader(path)
            return "\n".join(page.extract_text() for page in reader.pages)

        case "docx":
            with open(path, "rb") as file:
                source_stream = BytesIO(file.read())

            document = DocxDocument(source_stream)
            return "\n".join(paragraph.text for paragraph in document.paragraphs)

        case "txt" | "md":
            with open(path, encoding="utf-8") as file:
                return file.read()

        case "csv":
            with open(path, newline="", encoding="utf-8") as file:
                # Row-wise text, not a DataFrame: RAG over tabular data is
                # a different retrieval problem, and this is enough for
                # chunking/embedding a CSV as plain text.
                return "\n".join(", ".join(row) for row in csv.reader(file))

        case "xlsx":
            # read_only + data_only: stream cells instead of loading the
            # whole workbook, and return computed values rather than
            # formula strings.
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for sheet in workbook.worksheets:
                rows = [
                    ", ".join(str(cell) for cell in row if cell is not None)
                    for row in sheet.iter_rows(values_only=True)
                ]
                sheets.append(f"# {sheet.title}\n" + "\n".join(rows))
            workbook.close()
            return "\n\n".join(sheets)

        case _:
            raise UnsupportedFileTypeError(f"Unknown document type: {content_type}")
